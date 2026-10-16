"""役ごとの字幕 (文字色・話者名) のテスト"""

from pptx import Presentation
from pptx.dml.color import RGBColor

from pptx_writer import _add_subtitle_shapes
from tts.base import Timing


def runs(slide, shape_id):
    shape = next(s for s in slide.shapes if s.shape_id == shape_id)
    return [(r.text, str(r.font.color.rgb), bool(r.font.bold))
            for p in shape.text_frame.paragraphs for r in p.runs]


def make(timings, **kw):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ids, _ = _add_subtitle_shapes(slide, timings, prs, font_color=RGBColor(0xFF, 0xFF, 0xFF), **kw)
    return [runs(slide, i) for i in ids]


ROLES = {"A": ("", "ずんだもん"), "B": ("FF7FBF", "めたん")}


def test_role_color_is_the_default_color():
    assert make([Timing("一", 0, 1, "A"), Timing("二<color=#00FF00>緑</color>", 1, 1, "B")], roles=ROLES) == [
        [("一", "FFFFFF", False)],
        [("二", "FF7FBF", False), ("緑", "00FF00", False)],  # <color> は役の色より優先
    ]


def test_show_speaker_name():
    result = make([Timing("一\n二", 0, 1, "B"), Timing("三", 1, 1, "誰"), Timing("四", 2, 1, None)],
                  roles=ROLES, show_name=True)
    assert result == [
        [("めたん: ", "FF7FBF", True), ("一", "FF7FBF", False), ("二", "FF7FBF", False)],  # 名前は最初の行だけ
        [("誰: ", "FFFFFF", True), ("三", "FFFFFF", False)],  # 定義されていない役は役名のまま
        [("四", "FFFFFF", False)],
    ]


def test_name_with_tag_characters_is_text():
    assert make([Timing("一", 0, 1, "A")], roles={"A": ("", "<b>A")}, show_name=True) == [
        [("<b>A: ", "FFFFFF", True), ("一", "FFFFFF", False)]]


def test_without_roles_nothing_changes():
    assert make([Timing("一", 0, 1, "A")], show_name=False) == [[("一", "FFFFFF", False)]]
